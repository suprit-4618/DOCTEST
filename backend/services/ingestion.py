import io
import json
import logging
import os
import re
import shutil
import subprocess
import zipfile
from pathlib import Path
from typing import List, Optional, Tuple

import fitz  # PyMuPDF
from PIL import Image, ImageDraw, ImageFont, ImageOps, ImageEnhance
import docx

from app.core.config import settings
from app.core.database import SessionLocal
from models.db import DocumentDB, QuestionDB
from models.enums import DocumentStatus, FileType, AnswerSource

logger = logging.getLogger(__name__)


class IngestionError(Exception):
    """Custom exception raised during document validation or processing."""
    pass


def sanitize_filename(filename: str) -> str:
    """
    Sanitizes user-provided filenames to prevent path traversal,
    directory escapes, or shell injection vulnerabilities.
    """
    if not filename:
        return "document"
    # Strip any directory components (cross-platform)
    base = os.path.basename(filename.replace("\\", "/"))
    # Replace non-alphanumeric (except dot, dash, underscore) with underscore
    clean = re.sub(r"[^a-zA-Z0-9_.-]", "_", base)
    # Strip leading/trailing dots or spaces to prevent hidden/empty files
    clean = clean.strip(" ._")
    return clean or "document"


def get_document_storage_dir(document_id: str) -> Path:
    """Returns and ensures the storage directory for a specific document."""
    # Ensure document_id is clean UUID / safe string
    safe_id = sanitize_filename(document_id)
    doc_dir = Path(settings.STORAGE_DIR) / safe_id
    doc_dir.mkdir(parents=True, exist_ok=True)
    return doc_dir


def validate_file_content_and_type(file_bytes: bytes, filename: str) -> Tuple[FileType, str]:
    """
    Validate file size and detect file type by inspecting magic bytes / content.
    Returns (FileType, normalized_extension).
    """
    # 1. Size Validation
    if len(file_bytes) > settings.MAX_FILE_SIZE_BYTES:
        max_mb = settings.MAX_FILE_SIZE_BYTES / (1024 * 1024)
        raise IngestionError(f"File size ({len(file_bytes) / (1024 * 1024):.1f} MB) exceeds limit of {max_mb:.0f} MB.")

    if len(file_bytes) == 0:
        raise IngestionError("Uploaded file is empty.")

    # 2. Content / Magic Byte Inspection
    # PDF check
    if file_bytes.startswith(b"%PDF"):
        return FileType.PDF, ".pdf"

    # PNG check
    if file_bytes.startswith(b"\x89PNG\r\n\x1a\n"):
        return FileType.IMAGE, ".png"

    # JPEG check
    if file_bytes.startswith(b"\xff\xd8\xff"):
        return FileType.IMAGE, ".jpg"

    # WEBP check
    if len(file_bytes) > 12 and file_bytes[:4] == b"RIFF" and file_bytes[8:12] == b"WEBP":
        return FileType.IMAGE, ".webp"

    # DOCX check (ZIP archive containing word/document.xml)
    if file_bytes.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
                namelist = zf.namelist()
                if "word/document.xml" in namelist or any(n.startswith("word/") for n in namelist):
                    return FileType.DOCX, ".docx"
        except zipfile.BadZipFile:
            pass

    # Other Image formats via Pillow validation
    try:
        with Image.open(io.BytesIO(file_bytes)) as img:
            img.verify()
            fmt = (img.format or "").lower()
            if fmt in ["png", "jpeg", "jpg", "webp", "bmp", "tiff", "heic"]:
                return FileType.IMAGE, f".{fmt}"
    except Exception:
        pass

    # Check if plain text
    try:
        file_bytes.decode("utf-8")
        ext = os.path.splitext(filename)[1].lower()
        if ext in [".txt", ".md", ".csv", ".json", ""]:
            return FileType.TEXT, ".txt"
    except UnicodeDecodeError:
        pass

    raise IngestionError(
        f"Unsupported or invalid file format for '{filename}'. Supported types: PDF, PNG, JPG/JPEG, WEBP, DOCX, or plain text."
    )


def normalize_image(
    image_bytes: bytes,
    max_dim: int = settings.IMAGE_MAX_DIMENSION,
    enhance_photo: bool = True
) -> Image.Image:
    """
    Normalizes an image:
    1. Auto-rotates using EXIF orientation metadata.
    2. Converts to standard RGB color mode.
    3. Enhances scanned/low-quality photos: auto-contrast (reduces shadows/uneven lighting) + subtle sharpening.
    4. Caches/caps longest dimension to max_dim with high-quality resampling.
    """
    img = Image.open(io.BytesIO(image_bytes))
    
    # 1. Auto-rotate based on EXIF
    try:
        img = ImageOps.exif_transpose(img)
    except Exception as e:
        logger.debug(f"EXIF transpose skipped or failed: {e}")

    # 2. Convert to RGB mode
    if img.mode not in ("RGB", "L"):
        img = img.convert("RGB")
    elif img.mode == "L":
        img = img.convert("RGB")

    # 3. Preprocessing for scanned/shadowed/low-quality photos
    if enhance_photo:
        try:
            # Auto-contrast balances uneven shadowing and washed-out scans
            img = ImageOps.autocontrast(img, cutoff=1)
            # Sharpening helps clarify blurry text characters
            enhancer = ImageEnhance.Sharpness(img)
            img = enhancer.enhance(1.25)
        except Exception as e:
            logger.debug(f"Image enhancement filter skipped: {e}")

    # 4. Resize if exceeding max dimension
    w, h = img.size
    if max(w, h) > max_dim:
        if w > h:
            new_w = max_dim
            new_h = int(h * (max_dim / w))
        else:
            new_h = max_dim
            new_w = int(w * (max_dim / h))
        img = img.resize((new_w, new_h), Image.Resampling.LANCZOS)

    return img


def render_text_to_images(text: str, dpi: int = 150) -> List[Image.Image]:
    """
    Renders text onto formatted A4-proportioned image pages.
    """
    # A4 standard at 150 DPI: ~1240 x 1754 px
    page_w, page_h = 1240, 1754
    margin = 80
    usable_w = page_w - (2 * margin)
    line_height = 30
    
    # Font setup
    try:
        font = ImageFont.truetype("arial.ttf", 20)
    except IOError:
        font = ImageFont.load_default()

    # Split text into lines
    raw_lines = text.splitlines()
    wrapped_lines: List[str] = []
    
    # Character estimate per line based on usable width
    approx_chars_per_line = int(usable_w / 11)
    for raw_line in raw_lines:
        if not raw_line.strip():
            wrapped_lines.append("")
            continue
        while len(raw_line) > approx_chars_per_line:
            split_idx = raw_line.rfind(" ", 0, approx_chars_per_line)
            if split_idx == -1:
                split_idx = approx_chars_per_line
            wrapped_lines.append(raw_line[:split_idx])
            raw_line = raw_line[split_idx:].lstrip()
        wrapped_lines.append(raw_line)

    lines_per_page = (page_h - (2 * margin)) // line_height
    pages: List[Image.Image] = []
    
    total_lines = len(wrapped_lines)
    for p_idx in range(0, max(1, total_lines), max(1, lines_per_page)):
        page_lines = wrapped_lines[p_idx: p_idx + lines_per_page]
        img = Image.new("RGB", (page_w, page_h), color=(255, 255, 255))
        draw = ImageDraw.Draw(img)
        
        y = margin
        for line in page_lines:
            draw.text((margin, y), line, fill=(20, 20, 20), font=font)
            y += line_height
        pages.append(img)

    return pages


def convert_docx_with_libreoffice(docx_path: Path, output_dir: Path) -> Optional[Path]:
    """
    Attempts to convert a DOCX file to PDF using headless LibreOffice if installed.
    """
    soffice_candidates = [
        "soffice",
        "libreoffice",
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    soffice_bin = None
    for cand in soffice_candidates:
        if shutil.which(cand) or os.path.exists(cand):
            soffice_bin = cand
            break

    if not soffice_bin:
        return None

    try:
        cmd = [
            soffice_bin,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_dir),
            str(docx_path)
        ]
        result = subprocess.run(cmd, capture_output=True, text=True, timeout=30)
        if result.returncode == 0:
            pdf_path = output_dir / f"{docx_path.stem}.pdf"
            if pdf_path.exists():
                return pdf_path
    except Exception as e:
        logger.warning(f"LibreOffice conversion failed: {e}")

    return None


def extract_docx_text(file_bytes: bytes) -> str:
    """Extracts raw text content from DOCX file bytes using python-docx."""
    doc = docx.Document(io.BytesIO(file_bytes))
    full_text = []
    for para in doc.paragraphs:
        if para.text:
            full_text.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                full_text.append(" | ".join(row_text))
    return "\n\n".join(full_text)


def process_document_sync(
    document_id: str,
    file_bytes: Optional[bytes] = None,
    file_type: Optional[FileType] = None,
    original_filename: Optional[str] = None,
    pasted_text: Optional[str] = None,
) -> int:
    """
    Core synchronous document ingestion logic:
    Converts input into a list of page images stored under /storage/{document_id}/page_{n}.png.
    Returns total page count.
    """
    doc_dir = get_document_storage_dir(document_id)
    db = SessionLocal()

    try:
        doc = db.query(DocumentDB).filter(DocumentDB.id == document_id).first()
        if not doc:
            raise IngestionError(f"Document {document_id} not found in database.")

        doc.status = DocumentStatus.PROCESSING.value
        doc.error_message = None
        db.commit()

        rendered_pages: List[str] = []

        # Case 1: Pasted Text
        if pasted_text is not None or file_type == FileType.TEXT:
            text_content = pasted_text or (file_bytes.decode("utf-8", errors="ignore") if file_bytes else "")
            doc.extracted_text = text_content
            
            # Save original text file
            txt_file = doc_dir / "original.txt"
            txt_file.write_text(text_content, encoding="utf-8")

            # Render text to visual page images
            page_images = render_text_to_images(text_content, dpi=settings.PDF_RENDER_DPI)
            total_pages = len(page_images)
            doc.page_count = total_pages
            db.commit()

            for i, img in enumerate(page_images):
                page_filename = f"page_{i + 1}.png"
                img_path = doc_dir / page_filename
                img.save(img_path, format="PNG")
                rendered_pages.append(page_filename)
                
                doc.pages_processed = i + 1
                db.commit()

        # Case 2: PDF Document
        elif file_type == FileType.PDF:
            assert file_bytes is not None
            # Save original PDF
            pdf_file = doc_dir / "original.pdf"
            pdf_file.write_bytes(file_bytes)

            fitz_doc = fitz.open(stream=file_bytes, filetype="pdf")
            total_pages = fitz_doc.page_count

            if total_pages > settings.MAX_PAGE_COUNT:
                raise IngestionError(
                    f"PDF has {total_pages} pages, which exceeds the maximum limit of {settings.MAX_PAGE_COUNT} pages."
                )

            doc.page_count = total_pages
            db.commit()

            # Render each page to PNG at ~150 DPI (zoom 150/72)
            zoom = settings.PDF_RENDER_DPI / 72.0
            matrix = fitz.Matrix(zoom, zoom)

            for i in range(total_pages):
                page = fitz_doc.load_page(i)
                pix = page.get_pixmap(matrix=matrix, alpha=False)
                page_filename = f"page_{i + 1}.png"
                pix.save(str(doc_dir / page_filename))
                rendered_pages.append(page_filename)

                doc.pages_processed = i + 1
                db.commit()

            fitz_doc.close()

        # Case 3: Image File (PNG, JPG, WEBP, etc.)
        elif file_type == FileType.IMAGE:
            assert file_bytes is not None
            ext = os.path.splitext(original_filename or "image.png")[1] or ".png"
            orig_file = doc_dir / f"original{ext}"
            orig_file.write_bytes(file_bytes)

            img = normalize_image(file_bytes, max_dim=settings.IMAGE_MAX_DIMENSION)
            page_filename = "page_1.png"
            img.save(doc_dir / page_filename, format="PNG")
            rendered_pages.append(page_filename)

            doc.page_count = 1
            doc.pages_processed = 1
            db.commit()

        # Case 4: DOCX Document
        elif file_type == FileType.DOCX:
            assert file_bytes is not None
            docx_file = doc_dir / "original.docx"
            docx_file.write_bytes(file_bytes)

            # Try LibreOffice conversion first
            converted_pdf = convert_docx_with_libreoffice(docx_file, doc_dir)
            if converted_pdf and converted_pdf.exists():
                logger.info(f"Successfully converted DOCX to PDF with LibreOffice: {converted_pdf}")
                pdf_bytes = converted_pdf.read_bytes()
                fitz_doc = fitz.open(stream=pdf_bytes, filetype="pdf")
                total_pages = fitz_doc.page_count

                if total_pages > settings.MAX_PAGE_COUNT:
                    raise IngestionError(f"DOCX contains {total_pages} pages, exceeding {settings.MAX_PAGE_COUNT} page limit.")

                doc.page_count = total_pages
                db.commit()

                zoom = settings.PDF_RENDER_DPI / 72.0
                matrix = fitz.Matrix(zoom, zoom)
                for i in range(total_pages):
                    page = fitz_doc.load_page(i)
                    pix = page.get_pixmap(matrix=matrix, alpha=False)
                    page_filename = f"page_{i + 1}.png"
                    pix.save(str(doc_dir / page_filename))
                    rendered_pages.append(page_filename)

                    doc.pages_processed = i + 1
                    db.commit()
                fitz_doc.close()
            else:
                # Fallback: python-docx text extraction
                logger.warning(
                    "LibreOffice not available for DOCX rendering; falling back to python-docx text extraction. "
                    "Highlight information and rich formatting may be lost."
                )
                raw_text = extract_docx_text(file_bytes)
                doc.extracted_text = raw_text
                
                # Render text to visual page images
                page_images = render_text_to_images(raw_text, dpi=settings.PDF_RENDER_DPI)
                total_pages = max(1, len(page_images))
                doc.page_count = total_pages
                db.commit()

                for i, img in enumerate(page_images):
                    page_filename = f"page_{i + 1}.png"
                    img.save(doc_dir / page_filename, format="PNG")
                    rendered_pages.append(page_filename)

                    doc.pages_processed = i + 1
                    db.commit()

        else:
            raise IngestionError(f"Unknown or unhandled file type: {file_type}")

        # Save manifest.json
        manifest = {
            "document_id": document_id,
            "filename": doc.filename,
            "file_type": doc.file_type,
            "total_pages": len(rendered_pages),
            "pages": rendered_pages,
        }
        (doc_dir / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

        # If raw text or pasted text is present, extract questions from it directly
        if doc.extracted_text:
            from services.extraction import parse_text_questions
            text_extracted_qs = parse_text_questions(doc.extracted_text)
            if text_extracted_qs:
                db.query(QuestionDB).filter(QuestionDB.document_id == doc.id).delete()
                for idx, tq in enumerate(text_extracted_qs):
                    new_q = QuestionDB(
                        document_id=doc.id,
                        number=tq.question_number or (idx + 1),
                        text=tq.question_text,
                        options=[opt.model_dump() for opt in tq.options],
                        correct_options=tq.correct_options,
                        answer_source=tq.answer_source,
                        confidence=tq.confidence,
                        source_page=tq.source_page or 1,
                        explanation=tq.explanation,
                        needs_review=tq.needs_review
                    )
                    db.add(new_q)
                doc.status = DocumentStatus.EXTRACTED.value
                db.commit()

        # Update processing progress
        doc.pages_processed = doc.page_count
        has_api_key = bool(settings.GEMINI_API_KEY or settings.OPENAI_API_KEY)
        q_count = db.query(QuestionDB).filter(QuestionDB.document_id == document_id).count()
        if q_count > 0 or not has_api_key:
            doc.status = DocumentStatus.EXTRACTED.value
        else:
            doc.status = DocumentStatus.PROCESSING.value
        db.commit()

        logger.info(f"Document {document_id} ingestion completed successfully with {doc.page_count} pages.")
        return doc.page_count

    except Exception as e:
        logger.error(f"Ingestion failed for document {document_id}: {str(e)}", exc_info=True)
        doc = db.query(DocumentDB).filter(DocumentDB.id == document_id).first()
        if doc:
            doc.status = DocumentStatus.FAILED.value
            doc.error_message = str(e)
            db.commit()
        raise
    finally:
        db.close()


async def process_document_background(
    document_id: str,
    file_bytes: Optional[bytes] = None,
    file_type: Optional[FileType] = None,
    original_filename: Optional[str] = None,
    pasted_text: Optional[str] = None,
):
    """
    Asynchronous background worker:
    1. Renders document into high-resolution page images.
    2. If an AI API Key (Gemini/OpenAI) is configured, runs the Vision LLM Extraction Pipeline automatically.
    3. If no key is set, completes ingestion with extracted text / 0 questions so user can review or set key.
    """
    db = SessionLocal()
    try:
        # Step 1: Ingest & render page images
        process_document_sync(
            document_id=document_id,
            file_bytes=file_bytes,
            file_type=file_type,
            original_filename=original_filename,
            pasted_text=pasted_text
        )

        doc = db.query(DocumentDB).filter(DocumentDB.id == document_id).first()
        if not doc:
            return

        # Check if AI Vision extraction is available
        has_api_key = bool(settings.GEMINI_API_KEY or settings.OPENAI_API_KEY)
        
        # If questions were already parsed from pasted text or DOCX, mark EXTRACTED
        existing_q_count = db.query(QuestionDB).filter(QuestionDB.document_id == document_id).count()
        if existing_q_count > 0:
            doc.status = DocumentStatus.EXTRACTED.value
            db.commit()
            return

        # If API key is configured and questions not yet parsed, run Vision LLM extraction
        if has_api_key:
            from services.extraction import ExtractionPipeline
            pipeline = ExtractionPipeline()
            await pipeline.run_extraction(document_id)
        else:
            # Complete ingestion gracefully without fake dummy questions
            doc.status = DocumentStatus.EXTRACTED.value
            doc.error_message = None
            db.commit()
            logger.info(
                f"Document {document_id} rendered {doc.page_count} pages. "
                "Set GEMINI_API_KEY or OPENAI_API_KEY in .env to enable automated AI Vision parsing."
            )

    except Exception as e:
        logger.error(f"Background extraction failed for document {document_id}: {e}", exc_info=True)
        doc = db.query(DocumentDB).filter(DocumentDB.id == document_id).first()
        if doc and not doc.is_cancelled:
            doc.status = DocumentStatus.FAILED.value
            doc.error_message = f"Extraction error: {str(e)}"
            db.commit()
    finally:
        db.close()
